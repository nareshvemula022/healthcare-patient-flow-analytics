"""Phase 3 checks for create_excel_report.py and the workbook it creates.

The workbook is built from a small synthetic database in a temporary folder.
"""

import json
import re
import shutil
import sqlite3
import subprocess
import zipfile
from contextlib import closing
from datetime import datetime

import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart
from openpyxl.utils import get_column_letter

import analyze_data
import config
import create_excel_report
from conftest import PROJECT_ROOT
from create_excel_report import ReportError

SHEETS = ["Summary", "Monthly Trends", "Department Analysis", "Provider Analysis",
          "Appointment Types", "Data Quality", "Appointment Data", "Capacity Data"]
ANALYSIS_SHEETS = {  # sheet -> name of the SQL result it shows
    "Monthly Trends": "monthly_trends",
    "Department Analysis": "department_performance",
    "Provider Analysis": "provider_performance",
    "Appointment Types": "appointment_type_performance",
}
HEADER_ROW = 4
RATE_PATTERN = re.compile(r'^=IF\(([A-Z]+)(\d+)=0,"",([A-Z]+)(\d+)/([A-Z]+)(\d+)\)$')


@pytest.fixture(scope="session")
def workbook_path(prepared_data, tmp_path_factory):
    path = tmp_path_factory.mktemp("report") / "report.xlsx"
    create_excel_report.build_report(prepared_data["database"], prepared_data["quality_report"],
                                     path)
    return path


@pytest.fixture(scope="session")
def workbook(workbook_path):
    return load_workbook(workbook_path)


@pytest.fixture(scope="session")
def sql_results(prepared_data):
    results, _checks = analyze_data.load_analytics(prepared_data["database"])
    return results


def headers(ws, row=HEADER_ROW):
    """Header text -> column letter."""
    return {cell.value: cell.column_letter for cell in ws[row] if cell.value}


def kpi_rows(ws):
    """Summary KPI name -> row number."""
    return {ws.cell(row=row, column=1).value: row for row in range(1, ws.max_row + 1)
            if ws.cell(row=row, column=1).value in analyze_data.KPI_DEFINITIONS}


def expected_rate(numerator_col, denominator_col, row):
    return f'=IF({denominator_col}{row}=0,"",{numerator_col}{row}/{denominator_col}{row})'


def evaluate_rate(ws, formula):
    """Work out =IF(D=0,"",N/D) from the cell values, the way Excel would."""
    _d_col, _d_row, n_col, n_row, d_col, d_row = RATE_PATTERN.match(formula).groups()
    denominator = ws[f"{d_col}{d_row}"].value
    return None if denominator == 0 else ws[f"{n_col}{n_row}"].value / denominator


# --- Structure ------------------------------------------------------------------------

def test_workbook_opens_with_sheets_in_order(workbook):
    assert workbook.sheetnames == SHEETS


def test_workbook_recalculates_automatically(workbook):
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.fullCalcOnLoad is True


def test_no_external_links_or_macros(workbook, workbook_path):
    assert workbook._external_links == []
    names = zipfile.ZipFile(workbook_path).namelist()
    assert not any("externalLink" in name or "vbaProject" in name for name in names)


def test_no_hidden_sheets(workbook):
    assert all(ws.sheet_state == "visible" for ws in workbook.worksheets)


CHART_TITLES = ["Monthly Appointment Volume (all statuses)", "Appointments by Status",
                "Appointments by Department (all statuses)"]
EMU_PER_PIXEL = 9525


def chart_title(chart):
    return "".join(run.t for paragraph in chart.title.tx.rich.p for run in paragraph.r)


def summary_charts(workbook):
    """Summary charts by title."""
    return {chart_title(chart): chart for chart in workbook["Summary"]._charts}


def column_pixels(ws, column):
    width = ws.column_dimensions[get_column_letter(column)].width or 8.43
    return int(width * 7 + 5)   # Excel's pixel width for Calibri 11


def row_pixels(ws, row):
    return (ws.row_dimensions[row].height or 15) * 96 / 72


def chart_box(ws, chart):
    """(left, top, right, bottom) of a chart on the sheet, in pixels."""
    start = chart.anchor._from
    left = sum(column_pixels(ws, c) for c in range(1, start.col + 1)) + start.colOff / EMU_PER_PIXEL
    top = sum(row_pixels(ws, r) for r in range(1, start.row + 1)) + start.rowOff / EMU_PER_PIXEL
    return (left, top, left + chart.anchor.ext.width / EMU_PER_PIXEL,
            top + chart.anchor.ext.height / EMU_PER_PIXEL)


def test_summary_has_three_bar_charts_and_no_pie(workbook):
    charts = workbook["Summary"]._charts
    assert len(charts) == 3
    assert all(isinstance(chart, BarChart) for chart in charts)
    assert not any(isinstance(chart, PieChart) for chart in charts)


def test_summary_chart_titles(workbook):
    assert sorted(summary_charts(workbook)) == sorted(CHART_TITLES)


def test_summary_charts_have_no_legend_or_axis_titles(workbook):
    for title, chart in summary_charts(workbook).items():
        assert chart.legend is None, title           # no "Series1" legend
        assert chart.x_axis.title is None, title
        assert chart.y_axis.title is None, title     # nothing crowding the tick labels
        assert chart.x_axis.delete is False and chart.y_axis.delete is False, title


def test_status_chart_uses_short_labels_without_overlapping_text(workbook):
    chart = summary_charts(workbook)["Appointments by Status"]
    assert chart.type == "col"
    assert len(chart.series) == 1
    series = chart.series[0]
    assert [point.v for point in series.cat.strLit.pt] == ["Completed", "Cancelled", "No Show"]
    assert series.tx.v == "Appointments"
    labels = chart.dataLabels
    assert labels.showVal is True
    assert labels.showCatName is False and labels.showSerName is False
    assert labels.showPercent is False and labels.showLegendKey is False


def test_status_chart_values_are_the_status_kpis(workbook):
    ws = workbook["Summary"]
    rows = kpi_rows(ws)
    chart = summary_charts(workbook)["Appointments by Status"]
    reference = chart.series[0].val.numRef.f
    assert reference == (f"'Summary'!$B${rows['Completed Appointments']}"
                         f":$B${rows['No Show Appointments']}")


def test_department_chart_value_axis_at_bottom(workbook):
    chart = summary_charts(workbook)["Appointments by Department (all statuses)"]
    assert chart.type == "bar"
    assert chart.x_axis.scaling.orientation == "maxMin"   # first department at the top
    assert chart.y_axis.crosses == "max"                  # value axis at the bottom


def test_monthly_chart_uses_short_month_labels(workbook):
    chart = summary_charts(workbook)["Monthly Appointment Volume (all statuses)"]
    assert chart.x_axis.numFmt.formatCode == "mmm"


def test_summary_charts_do_not_overlap(workbook):
    ws = workbook["Summary"]
    boxes = {title: chart_box(ws, chart) for title, chart in summary_charts(workbook).items()}
    titles = list(boxes)
    for i, first in enumerate(titles):
        for second in titles[i + 1:]:
            a, b = boxes[first], boxes[second]
            horizontal_gap = max(b[0] - a[2], a[0] - b[2])
            vertical_gap = max(b[1] - a[3], a[1] - b[3])
            assert max(horizontal_gap, vertical_gap) >= 10, (first, second)

    # Charts sit below the KPI table and within columns A-C.
    kpi_bottom = sum(row_pixels(ws, r) for r in range(1, max(kpi_rows(ws).values()) + 2))
    table_right = sum(column_pixels(ws, c) for c in range(1, 4))
    for title, (left, top, right, _bottom) in boxes.items():
        assert top > kpi_bottom, title
        assert right <= table_right, title


def test_summary_states_data_is_synthetic(workbook):
    text = " ".join(str(cell.value) for row in workbook["Summary"]["A1:A4"] for cell in row)
    assert "SYNTHETIC" in text.upper()


# --- Summary values and formulas -----------------------------------------------------------

def test_headline_kpi_values_match_sql(workbook, sql_results):
    ws = workbook["Summary"]
    rows = kpi_rows(ws)
    overall = sql_results["overall_kpis"].iloc[0]
    for name, source in [("Total Appointments", "total_appointments"),
                         ("Unique Patients", "unique_patients"),
                         ("Completed Appointments", "completed_appointments"),
                         ("Cancelled Appointments", "cancelled_appointments"),
                         ("No Show Appointments", "no_show_appointments"),
                         ("Available Slots", "available_slots")]:
        assert ws[f"B{rows[name]}"].value == overall[source], name
        assert ws[f"B{rows[name]}"].number_format == "#,##0"
    for name, source in [("Average Wait Minutes", "avg_wait_minutes"),
                         ("Average Visit Minutes", "avg_visit_minutes")]:
        assert ws[f"B{rows[name]}"].value == pytest.approx(overall[source])
        assert ws[f"B{rows[name]}"].number_format == "0.0"


def test_summary_rate_formulas(workbook, sql_results):
    ws = workbook["Summary"]
    rows = kpi_rows(ws)
    overall = sql_results["overall_kpis"].iloc[0]
    for name, numerator, denominator, source in [
        ("Completion Rate", "Completed Appointments", "Total Appointments", "completion_rate"),
        ("Cancellation Rate", "Cancelled Appointments", "Total Appointments",
         "cancellation_rate"),
        ("No Show Rate", "No Show Appointments", "Total Appointments", "no_show_rate"),
        ("Booking Rate", "Total Appointments", "Available Slots", "booking_rate"),
        ("Completed Slot Utilization", "Completed Appointments", "Available Slots",
         "completed_slot_utilization"),
    ]:
        cell = ws[f"B{rows[name]}"]
        assert cell.value == (f'=IF(B{rows[denominator]}=0,"",'
                              f'B{rows[numerator]}/B{rows[denominator]})'), name
        assert cell.number_format == "0.0%"
        assert evaluate_rate(ws, cell.value) == pytest.approx(overall[source]), name


# --- Analysis sheets --------------------------------------------------------------------------

@pytest.mark.parametrize("sheet", ANALYSIS_SHEETS)
def test_analysis_rate_formulas_reference_their_own_row(workbook, sql_results, sheet):
    ws = workbook[sheet]
    columns = headers(ws)
    expected = sql_results[ANALYSIS_SHEETS[sheet]]
    first_row = HEADER_ROW + 1
    rates = [("Completion Rate", "Completed", "Total Appointments", "completion_rate"),
             ("Cancellation Rate", "Cancelled", "Total Appointments", "cancellation_rate"),
             ("No Show Rate", "No Show", "Total Appointments", "no_show_rate")]
    if "Available Slots" in columns:
        rates += [("Booking Rate", "Total Appointments", "Available Slots", "booking_rate"),
                  ("Completed Slot Utilization", "Completed", "Available Slots",
                   "completed_slot_utilization")]

    for offset, sql_row in enumerate(expected.to_dict("records")):
        row = first_row + offset
        for rate, numerator, denominator, source in rates:
            cell = ws[f"{columns[rate]}{row}"]
            assert cell.value == expected_rate(columns[numerator], columns[denominator], row)
            assert cell.number_format == "0.0%"
            assert evaluate_rate(ws, cell.value) == pytest.approx(sql_row[source])


@pytest.mark.parametrize("sheet", ANALYSIS_SHEETS)
def test_analysis_values_and_total_row(workbook, sql_results, sheet):
    ws = workbook[sheet]
    columns = headers(ws)
    expected = sql_results[ANALYSIS_SHEETS[sheet]]
    first_row, last_row = HEADER_ROW + 1, HEADER_ROW + len(expected)
    total_row = last_row + 1
    overall = sql_results["overall_kpis"].iloc[0]

    # Counts and durations are the SQL values.
    totals = [ws[f"{columns['Total Appointments']}{row}"].value
              for row in range(first_row, last_row + 1)]
    assert totals == expected["total_appointments"].tolist()
    waits = [ws[f"{columns['Avg Wait (min)']}{row}"].value
             for row in range(first_row, last_row + 1)]
    assert waits == pytest.approx(expected["avg_wait_minutes"].tolist())
    assert ws[f"{columns['Avg Wait (min)']}{first_row}"].number_format == "0.0"
    assert ws[f"{columns['Completed']}{first_row}"].number_format == "#,##0"

    # Total row: SUM of the column, and rates from the summed counts.
    assert ws[f"A{total_row}"].value == "Total"
    letter = columns["Total Appointments"]
    assert ws[f"{letter}{total_row}"].value == f"=SUM({letter}{first_row}:{letter}{last_row})"
    assert sum(totals) == overall["total_appointments"]
    completion = ws[f"{columns['Completion Rate']}{total_row}"].value
    assert completion == expected_rate(columns["Completed"], letter, total_row)


def test_monthly_rows_are_in_date_order(workbook):
    ws = workbook["Monthly Trends"]
    months = [ws[f"A{row}"].value for row in range(5, 17)]
    assert all(isinstance(month, datetime) for month in months)
    assert months == sorted(months)
    assert [m.month for m in months] == list(range(1, 13))
    assert ws["A5"].number_format == "mmm yyyy"


def test_appointment_types_have_no_capacity_columns(workbook):
    columns = headers(workbook["Appointment Types"])
    assert "Available Slots" not in columns and "Booking Rate" not in columns


def test_only_data_quality_mentions_data_quality(workbook):
    for ws in workbook.worksheets:
        if ws.title == "Data Quality":
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert "Data Quality" not in cell.value, (ws.title, cell.coordinate)


# --- Data Quality sheet ------------------------------------------------------------------------

def test_data_quality_sheet(workbook, prepared_data):
    ws = workbook["Data Quality"]
    values = {ws.cell(row=row, column=1).value: ws.cell(row=row, column=2).value
              for row in range(1, ws.max_row + 1)}
    assert values["Synthetic data"].startswith("Yes")
    assert values["Phase 2 pipeline status"] == "Passed"
    assert values["Phase 2 validation violations"] == 0
    assert values["SQLite foreign_key_check"] == "Passed (0 violations)"

    validation = ws.tables["tblValidationChecks"]
    rows = ws[validation.ref][1:]
    assert len(rows) > 0 and all(row[2].value == 0 for row in rows)
    reconciliation = ws.tables["tblReconciliation"]
    assert all(row[0].value == "PASS" for row in ws[reconciliation.ref][1:])

    quality = json.loads(prepared_data["quality_report"].read_text(encoding="utf-8"))
    counts = {row[0].value: row[4].value for row in ws[ws.tables["tblRowCounts"].ref][1:]}
    assert counts == quality["database_row_counts"]


# --- Data sheets ---------------------------------------------------------------------------------

def test_data_rows_match_database(workbook, prepared_data):
    with closing(sqlite3.connect(prepared_data["database"])) as connection:
        appointments = connection.execute("SELECT COUNT(*) FROM appointments").fetchone()[0]
        schedules = connection.execute("SELECT COUNT(*) FROM provider_schedules").fetchone()[0]
    assert workbook["Appointment Data"].max_row - 1 == appointments
    assert workbook["Capacity Data"].max_row - 1 == schedules


@pytest.mark.parametrize("sheet, table", [
    ("Monthly Trends", "tblMonthlyTrends"), ("Department Analysis", "tblDepartments"),
    ("Provider Analysis", "tblProviders"), ("Appointment Types", "tblAppointmentTypes"),
    ("Appointment Data", "tblAppointmentData"), ("Capacity Data", "tblCapacityData"),
])
def test_tables_have_filters_and_banding(workbook, sheet, table):
    ws = workbook[sheet]
    assert table in ws.tables
    excel_table = ws.tables[table]
    assert excel_table.autoFilter is not None
    assert excel_table.autoFilter.ref == excel_table.ref
    assert excel_table.tableStyleInfo.showRowStripes


@pytest.mark.parametrize("sheet, cell", [
    ("Monthly Trends", "B5"), ("Department Analysis", "C5"), ("Provider Analysis", "C5"),
    ("Appointment Types", "B5"), ("Appointment Data", "A2"), ("Capacity Data", "A2"),
])
def test_freeze_panes(workbook, sheet, cell):
    assert workbook[sheet].freeze_panes == cell


def test_appointment_data_types_and_formats(workbook):
    ws = workbook["Appointment Data"]
    columns = headers(ws, row=1)
    assert ws[f"{columns['appointment_id']}2"].data_type == "s"
    assert ws[f"{columns['patient_id']}2"].number_format == "@"

    date_cell = ws[f"{columns['appointment_date']}2"]
    assert date_cell.is_date and date_cell.number_format == "mm/dd/yyyy"
    scheduled = ws[f"{columns['scheduled_datetime']}2"]
    assert scheduled.is_date and scheduled.number_format == "mm/dd/yyyy hh:mm"
    assert scheduled.value.hour >= 8

    for row in range(2, ws.max_row + 1):
        status = ws[f"{columns['appointment_status']}{row}"].value
        check_in = ws[f"{columns['check_in_datetime']}{row}"]
        wait = ws[f"{columns['wait_minutes']}{row}"].value
        if status == "Completed":
            assert isinstance(check_in.value, datetime) and wait is not None
        else:
            assert check_in.value is None and wait is None


def test_capacity_data_dates(workbook):
    ws = workbook["Capacity Data"]
    columns = headers(ws, row=1)
    cell = ws[f"{columns['schedule_date']}2"]
    assert cell.is_date and cell.number_format == "mm/dd/yyyy"
    assert ws[f"{columns['provider_id']}2"].data_type == "s"


FILTER_ROOM = 4   # characters of column width taken by a filter button, at least


def test_column_widths_fit_headers_and_filter_buttons(workbook):
    for ws in workbook.worksheets:
        header_row = 1 if ws.title in ("Appointment Data", "Capacity Data") else HEADER_ROW
        if ws.title in ("Summary", "Data Quality"):
            continue
        for cell in ws[header_row]:
            width = ws.column_dimensions[cell.column_letter].width
            longest_word = max(len(word) for word in str(cell.value).split())
            assert width >= longest_word + FILTER_ROOM, (ws.title, cell.value, width)


@pytest.mark.parametrize("sheet", [*ANALYSIS_SHEETS, "Data Quality"])
def test_initial_zoom(workbook, sheet):
    assert workbook[sheet].sheet_view.zoomScale == 90


@pytest.mark.parametrize("sheet", ANALYSIS_SHEETS)
def test_analysis_headers_wrap_and_have_room(workbook, sheet):
    ws = workbook[sheet]
    assert ws.row_dimensions[HEADER_ROW].height >= 45   # three lines of header text
    for cell in ws[HEADER_ROW]:
        assert cell.alignment.wrap_text, cell.value
        assert cell.alignment.vertical == "center", cell.value


@pytest.mark.parametrize("sheet, header, minimum", [
    ("Monthly Trends", "Total Appointments", 17),
    ("Monthly Trends", "Cancellation Rate", 17),
    ("Monthly Trends", "Completed Slot Utilization", 16),
    ("Department Analysis", "Total Appointments", 17),
    ("Department Analysis", "Completed Slot Utilization", 16),
    ("Provider Analysis", "Total Appointments", 17),
    ("Provider Analysis", "Daily Slot Capacity", 13),
    ("Provider Analysis", "Completed Slot Utilization", 16),
    ("Appointment Types", "Share of Appointments", 17),
    ("Appointment Types", "Avg Visit (min)", 11),
])
def test_important_analysis_columns_are_wide_enough(workbook, sheet, header, minimum):
    ws = workbook[sheet]
    assert ws.column_dimensions[headers(ws)[header]].width >= minimum


def test_data_quality_count_columns_are_wide_enough(workbook):
    ws = workbook["Data Quality"]
    table = ws.tables["tblRowCounts"]
    header_row = ws[table.ref][0]
    for cell in header_row[2:]:   # Raw CSV Rows, Processed CSV Rows, Database Rows
        width = ws.column_dimensions[cell.column_letter].width
        assert width >= 20, cell.value
        # Every header line fits next to the filter button.
        lines = -(-len(cell.value) // (width - FILTER_ROOM - 1))
        assert ws.row_dimensions[header_row[0].row].height >= 15 * lines, cell.value


def test_data_quality_wrapped_rows_are_tall_enough(workbook):
    ws = workbook["Data Quality"]
    checked = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.alignment.wrap_text and isinstance(cell.value, str) and cell.row > 5:
                width = ws.column_dimensions[cell.column_letter].width
                lines = -(-len(cell.value) // (width - 4))
                height = ws.row_dimensions[cell.row].height or 15
                assert height >= 15 * lines, (cell.coordinate, cell.value)
                checked += 1
    assert checked > 0


# --- Inputs and Git ----------------------------------------------------------------------------

def test_missing_quality_report_gives_helpful_error(prepared_data, tmp_path):
    with pytest.raises(FileNotFoundError, match="prepare_data.py"):
        create_excel_report.build_report(prepared_data["database"], tmp_path / "missing.json",
                                         tmp_path / "report.xlsx")


def test_out_of_date_quality_report_is_rejected(prepared_data, tmp_path):
    quality = json.loads(prepared_data["quality_report"].read_text(encoding="utf-8"))
    quality["database_row_counts"]["appointments"] += 1
    stale = tmp_path / "stale.json"
    stale.write_text(json.dumps(quality), encoding="utf-8")
    with pytest.raises(ReportError, match="prepare_data.py"):
        create_excel_report.build_report(prepared_data["database"], stale,
                                         tmp_path / "report.xlsx")


def test_missing_database_gives_helpful_error(tmp_path):
    with pytest.raises(FileNotFoundError, match="prepare_data.py"):
        create_excel_report.build_report(tmp_path / "missing.db", tmp_path / "q.json",
                                         tmp_path / "report.xlsx")


def test_report_workbook_is_not_ignored_by_git():
    if shutil.which("git") is None or not (PROJECT_ROOT / ".git").exists():
        pytest.skip("git repository not available")
    relative = config.EXCEL_REPORT_FILE.relative_to(PROJECT_ROOT).as_posix()
    result = subprocess.run(["git", "check-ignore", "-q", relative], cwd=PROJECT_ROOT)
    assert result.returncode == 1, f"{relative} is ignored by .gitignore"


def test_analytics_exports_are_ignored_by_git():
    if shutil.which("git") is None or not (PROJECT_ROOT / ".git").exists():
        pytest.skip("git repository not available")
    relative = (config.ANALYTICS_OUTPUT_DIR / "overall_kpis.csv").relative_to(PROJECT_ROOT)
    result = subprocess.run(["git", "check-ignore", "-q", relative.as_posix()], cwd=PROJECT_ROOT)
    assert result.returncode == 0
