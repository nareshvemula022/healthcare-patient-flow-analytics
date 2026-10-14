"""Phase 3: build the Excel operational report from the SYNTHETIC SQLite database.

All data is fictional. Run this after prepare_data.py (PowerShell):
    python create_excel_report.py

Creates reports/healthcare_patient_flow_report.xlsx with these sheets:
    Summary, Monthly Trends, Department Analysis, Provider Analysis,
    Appointment Types, Data Quality, Appointment Data, Capacity Data

Counts and average durations come from the SQL views (via analyze_data.py).
Rates are Excel formulas (numerator / denominator, blank when the denominator
is 0). Excel recalculates them when the workbook is opened.
"""

import argparse
import json
import math
import sqlite3
import sys
from contextlib import closing
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrData, StrVal
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo

import analyze_data
import config
from prepare_data import TABLES

SHEET_NAMES = ["Summary", "Monthly Trends", "Department Analysis", "Provider Analysis",
               "Appointment Types", "Data Quality", "Appointment Data", "Capacity Data"]
WORKBOOK_AUTHOR = "Healthcare Patient Flow Analytics"   # creator and lastModifiedBy

# --- Look and feel: restrained blue palette, dark text --------------------------
FONT_NAME = "Calibri"
TEXT_COLOR = "1F2937"
TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True, color=TEXT_COLOR)
SECTION_FONT = Font(name=FONT_NAME, size=12, bold=True, color=TEXT_COLOR)
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color=TEXT_COLOR)
BODY_FONT = Font(name=FONT_NAME, size=11, color=TEXT_COLOR)
BOLD_FONT = Font(name=FONT_NAME, size=11, bold=True, color=TEXT_COLOR)
NOTE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="4B5563")
HEADER_FILL = PatternFill("solid", fgColor="C9DAEA")
TOTAL_FILL = PatternFill("solid", fgColor="EEF3F8")
LIGHT_LINE = Side(style="thin", color="BFC9D4")
CELL_BORDER = Border(left=LIGHT_LINE, right=LIGHT_LINE, top=LIGHT_LINE, bottom=LIGHT_LINE)
HEADER_BORDER = Border(left=LIGHT_LINE, right=LIGHT_LINE, top=LIGHT_LINE,
                       bottom=Side(style="medium", color="44546A"))
TOTAL_BORDER = Border(left=LIGHT_LINE, right=LIGHT_LINE, bottom=LIGHT_LINE,
                      top=Side(style="medium", color="44546A"))
TABLE_STYLE = "TableStyleMedium2"   # Excel's standard blue table style: banded rows
SCALE_COLORS = ("FFFFFF", "9DC3E6")  # white -> light blue for conditional formatting
CHART_COLOR = "2E6E9E"

# --- Layout -------------------------------------------------------------------------
ANALYSIS_ZOOM = 90           # initial zoom (%) for the analysis and Data Quality sheets
ANALYSIS_HEADER_HEIGHT = 45  # points: room for three wrapped header lines
DATA_HEADER_HEIGHT = 32      # points: other header rows
FILTER_BUTTON_ROOM = 5       # extra column width (characters) for the filter button
SUMMARY_WIDTHS = [30, 14, 70]
# Summary charts (cm). Two rows: the monthly chart across columns A-C, then the
# status chart under columns A-B and the department chart under column C.
WIDE_CHART_WIDTH = 21.4
STATUS_CHART_WIDTH = 7.6     # ends ~0.8 cm before column C, leaving a gap
DEPARTMENT_CHART_WIDTH = 13.0
CHART_HEIGHT = 7.5
CHART_ROWS = 16              # rows from one chart row to the next (~1 row gap)
STATUS_LABELS = ["Completed", "Cancelled", "No Show"]

# --- Number formats ------------------------------------------------------------
COUNT = "#,##0"
PERCENT = "0.0%"
MINUTES = "0.0"
DATE = "mm/dd/yyyy"
DATETIME = "mm/dd/yyyy hh:mm"
MONTH = "mmm yyyy"
TEXT = "@"
GENERAL = "General"


class ReportError(Exception):
    """The inputs for the workbook are missing or out of date."""


# --- Analysis table columns ---------------------------------------------------------
# key:      short name, used by formulas to find a column
# source:   column in the SQL result (values written as-is)
# rate:     (numerator key, denominator key): an Excel formula instead of a value
# share_of: key of a column; formula = this row / the total row
# total:    what the Total row shows: "label", "sum", "rate", "weighted" or "share"

def column(key, header, width, number_format, source=None, rate=None, share_of=None,
           total=None, weight=None, highlight=False):
    return {"key": key, "header": header, "width": width, "format": number_format,
            "source": source, "rate": rate, "share_of": share_of, "total": total,
            "weight": weight, "highlight": highlight}


STATUS_COUNTS = [
    column("total", "Total Appointments", 14, COUNT, source="total_appointments", total="sum"),
    column("completed", "Completed", 12, COUNT, source="completed_appointments", total="sum"),
    column("cancelled", "Cancelled", 12, COUNT, source="cancelled_appointments", total="sum"),
    column("no_show", "No Show", 11, COUNT, source="no_show_appointments", total="sum"),
]
STATUS_RATES = [
    column("completion_rate", "Completion Rate", 12, PERCENT, rate=("completed", "total"),
           total="rate"),
    column("cancellation_rate", "Cancellation Rate", 13, PERCENT, rate=("cancelled", "total"),
           total="rate"),
    column("no_show_rate", "No Show Rate", 12, PERCENT, rate=("no_show", "total"),
           total="rate", highlight=True),
]
DURATIONS = [
    column("avg_wait", "Avg Wait (min)", 11, MINUTES, source="avg_wait_minutes",
           total="weighted", weight="completed", highlight=True),
    column("avg_visit", "Avg Visit (min)", 11, MINUTES, source="avg_visit_minutes",
           total="weighted", weight="completed"),
]
CAPACITY = [
    column("slots", "Available Slots", 12, COUNT, source="available_slots", total="sum"),
    column("booking_rate", "Booking Rate", 12, PERCENT, rate=("total", "slots"), total="rate"),
    column("utilization", "Completed Slot Utilization", 15, PERCENT,
           rate=("completed", "slots"), total="rate"),
]

MONTHLY_COLUMNS = ([column("month", "Month", 12, MONTH, source="month_start", total="label")]
                   + STATUS_COUNTS + STATUS_RATES + DURATIONS + CAPACITY)
DEPARTMENT_COLUMNS = ([column("department_id", "Department ID", 13, TEXT,
                              source="department_id", total="label"),
                       column("department_name", "Department", 18, TEXT,
                              source="department_name")]
                      + STATUS_COUNTS + STATUS_RATES + DURATIONS + CAPACITY)
PROVIDER_COLUMNS = ([column("provider_id", "Provider ID", 12, TEXT, source="provider_id",
                            total="label"),
                     column("provider_name", "Provider", 24, TEXT, source="provider_name"),
                     column("department_name", "Department", 18, TEXT,
                            source="department_name"),
                     column("specialty", "Specialty", 27, TEXT, source="specialty"),
                     column("daily_capacity", "Daily Slot Capacity", 12, COUNT,
                            source="daily_slot_capacity")]
                    + STATUS_COUNTS + STATUS_RATES + DURATIONS + CAPACITY)
TYPE_COLUMNS = ([column("appointment_type", "Appointment Type", 18, TEXT,
                        source="appointment_type", total="label"),
                 STATUS_COUNTS[0],
                 column("share", "Share of Appointments", 13, PERCENT, share_of="total",
                        total="share")]
                + STATUS_COUNTS[1:] + STATUS_RATES + DURATIONS)

# Summary sheet: (KPI name, source column in vw_overall_kpis or a rate formula, format).
# Rates reference the count cells above them by KPI name.
SUMMARY_KPIS = [
    ("Total Appointments", "total_appointments", COUNT),
    ("Unique Patients", "unique_patients", COUNT),
    ("Completed Appointments", "completed_appointments", COUNT),
    ("Cancelled Appointments", "cancelled_appointments", COUNT),
    ("No Show Appointments", "no_show_appointments", COUNT),
    ("Completion Rate", ("Completed Appointments", "Total Appointments"), PERCENT),
    ("Cancellation Rate", ("Cancelled Appointments", "Total Appointments"), PERCENT),
    ("No Show Rate", ("No Show Appointments", "Total Appointments"), PERCENT),
    ("Average Wait Minutes", "avg_wait_minutes", MINUTES),
    ("Average Visit Minutes", "avg_visit_minutes", MINUTES),
    ("Available Slots", "available_slots", COUNT),
    ("Booking Rate", ("Total Appointments", "Available Slots"), PERCENT),
    ("Completed Slot Utilization", ("Completed Appointments", "Available Slots"), PERCENT),
]
SUMMARY_HEADER_ROW = 6


# --- Small helpers -------------------------------------------------------------------

def clean(value):
    """Turn NaN into an empty cell; leave everything else unchanged."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    return value


def rate_formula(numerator_cell, denominator_cell):
    """Numerator / denominator, or blank when the denominator is 0."""
    return f'=IF({denominator_cell}=0,"",{numerator_cell}/{denominator_cell})'


def us_date(iso_text):
    return date.fromisoformat(iso_text).strftime("%m/%d/%Y")


def write_title(ws, title, notes):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 26
    for row, note in enumerate(notes, start=2):
        ws.cell(row=row, column=1, value=note).font = NOTE_FONT


def write_section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT
    ws.row_dimensions[row].height = 20


def style_header(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = HEADER_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(ws, header_row, headers, rows, formats, table_name=None, wrap_columns=(),
                header_height=DATA_HEADER_HEIGHT):
    """Write a header and rows with light borders. With table_name, the range becomes
    an Excel table (banded rows and filter buttons). Returns the last data row."""
    for col, header in enumerate(headers, start=1):
        style_header(ws.cell(row=header_row, column=col, value=header))
    ws.row_dimensions[header_row].height = header_height

    for offset, values in enumerate(rows, start=1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=header_row + offset, column=col, value=clean(value))
            cell.font = BODY_FONT
            cell.border = CELL_BORDER
            cell.number_format = formats[col - 1]
            if col in wrap_columns:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    last_row = header_row + len(rows)
    if table_name and rows:
        ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
        table.autoFilter = AutoFilter(ref=ref)
        ws.add_table(table)
    return last_row


def set_widths(ws, widths):
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def header_width(header, width):
    """At least wide enough for the longest header word plus the filter button."""
    longest_word = max(len(word) for word in header.split())
    return max(width, longest_word + FILTER_BUTTON_ROOM)


def wrapped_lines(text, width):
    """Lines needed for wrapped text in a column (about one character per width unit,
    with some spare room for wide letters)."""
    return max(1, math.ceil(len(str(text or "")) / (width - 6)))


def fit_wrapped_rows(ws, first_row, last_row, col, width):
    """Make rows tall enough for wrapped text (openpyxl does not do this itself)."""
    for row in range(first_row, last_row + 1):
        lines = wrapped_lines(ws.cell(row=row, column=col).value, width)
        if lines > 1:
            ws.row_dimensions[row].height = 15 * lines + 3


# --- Analysis sheets -------------------------------------------------------------------

def write_analysis_sheet(ws, title, notes, df, columns, table_name, freeze):
    """One analysis table: SQL values, rate formulas, then a Total row underneath.
    Returns where things are, so charts and tests can find them."""
    header_row = len(notes) + 3
    first_row = header_row + 1
    last_row = header_row + len(df)
    total_row = last_row + 1
    letters = {c["key"]: get_column_letter(i) for i, c in enumerate(columns, start=1)}

    rows = []
    for row, record in enumerate(df.to_dict("records"), start=first_row):
        values = []
        for c in columns:
            if c["rate"]:
                numerator, denominator = (letters[key] for key in c["rate"])
                values.append(rate_formula(f"{numerator}{row}", f"{denominator}{row}"))
            elif c["share_of"]:
                share = letters[c["share_of"]]
                values.append(rate_formula(f"{share}{row}", f"${share}${total_row}"))
            else:
                values.append(record[c["source"]])
        rows.append(values)

    write_title(ws, title, notes)
    write_table(ws, header_row, [c["header"] for c in columns], rows,
                [c["format"] for c in columns], table_name,
                header_height=ANALYSIS_HEADER_HEIGHT)

    # Total row: sums of counts, rates from the summed counts (never an average of
    # percentages), and averages weighted by completed appointments.
    for col, c in enumerate(columns, start=1):
        letter = get_column_letter(col)
        cells = f"{letter}{first_row}:{letter}{last_row}"
        value = None
        if c["total"] == "label":
            value = "Total"
        elif c["total"] == "sum":
            value = f"=SUM({cells})"
        elif c["total"] == "rate":
            numerator, denominator = (letters[key] for key in c["rate"])
            value = rate_formula(f"{numerator}{total_row}", f"{denominator}{total_row}")
        elif c["total"] == "share":
            share = letters[c["share_of"]]
            value = rate_formula(f"{share}{total_row}", f"${share}${total_row}")
        elif c["total"] == "weighted":
            weight = letters[c["weight"]]
            weights = f"{weight}{first_row}:{weight}{last_row}"
            value = (f'=IF(SUM({weights})=0,"",'
                     f'SUMPRODUCT({cells},{weights})/SUM({weights}))')
        cell = ws.cell(row=total_row, column=col, value=value)
        cell.font = BOLD_FONT
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.number_format = c["format"] if c["total"] != "label" else GENERAL

    # Conditional formatting, used sparingly: no-show rate and average wait only.
    for c in columns:
        if c["highlight"]:
            letter = letters[c["key"]]
            ws.conditional_formatting.add(
                f"{letter}{first_row}:{letter}{last_row}",
                ColorScaleRule(start_type="min", start_color=SCALE_COLORS[0],
                               end_type="max", end_color=SCALE_COLORS[1]))

    set_widths(ws, [header_width(c["header"], c["width"]) for c in columns])
    ws.freeze_panes = freeze
    ws.sheet_view.zoomScale = ANALYSIS_ZOOM
    return {"letters": letters, "header_row": header_row, "first_row": first_row,
            "last_row": last_row, "total_row": total_row}


# --- Summary sheet ----------------------------------------------------------------------

def write_summary(ws, overall, monthly, monthly_sheet, department, department_sheet):
    period = f"{us_date(overall['period_start'])} to {us_date(overall['period_end'])}"
    write_title(ws, "Healthcare Patient Flow & Operational Analytics Report", [
        f"Reporting period: {period}",
        "SYNTHETIC DATA: every record is computer-generated. Nothing in this workbook "
        "describes real patients, providers, hospitals or clinics.",
    ])
    ws["A3"].font = BOLD_FONT
    write_section(ws, SUMMARY_HEADER_ROW - 1, "Key Performance Indicators")

    kpi_rows = {name: SUMMARY_HEADER_ROW + i for i, (name, _s, _f) in
                enumerate(SUMMARY_KPIS, start=1)}
    rows = []
    for name, source, _number_format in SUMMARY_KPIS:
        if isinstance(source, tuple):
            numerator, denominator = (f"B{kpi_rows[kpi]}" for kpi in source)
            value = rate_formula(numerator, denominator)
        else:
            value = overall[source]
        rows.append([name, value, analyze_data.KPI_DEFINITIONS[name]])
    last_row = write_table(ws, SUMMARY_HEADER_ROW, ["KPI", "Value", "Definition"], rows,
                           [GENERAL, COUNT, GENERAL])
    for name, _source, number_format in SUMMARY_KPIS:
        ws[f"B{kpi_rows[name]}"].number_format = number_format
        ws[f"A{kpi_rows[name]}"].font = BOLD_FONT

    note_row = last_row + 1
    ws.cell(row=note_row, column=1, value=(
        "Counts and averages come from the SQL views in sql/02_create_analytics_views.sql. "
        "Rates are Excel formulas using the overall counts above.")).font = NOTE_FONT
    set_widths(ws, SUMMARY_WIDTHS)

    # Charts. Titles describe what is shown, nothing more. No legends: each chart has
    # one series, and its title says what the bars are.
    chart_row = note_row + 2
    write_section(ws, chart_row, "Charts")

    volume = BarChart()
    volume.type = "col"
    volume.title = "Monthly Appointment Volume (all statuses)"
    total_col = column_index(MONTHLY_COLUMNS, "total")
    volume.add_data(Reference(monthly_sheet, min_col=total_col, min_row=monthly["header_row"],
                              max_row=monthly["last_row"]), titles_from_data=True)
    volume.set_categories(Reference(monthly_sheet, min_col=1, min_row=monthly["first_row"],
                                    max_row=monthly["last_row"]))
    volume.x_axis.number_format = "mmm"   # short month labels; the year is in the period line
    style_chart(volume)
    volume.width, volume.height = WIDE_CHART_WIDTH, CHART_HEIGHT
    ws.add_chart(volume, f"A{chart_row + 1}")

    # Status counts come from the KPI cells above. The category labels are fixed short
    # words, so the chart does not show the longer KPI names.
    status = BarChart()
    status.type = "col"
    status.title = "Appointments by Status"
    status.add_data(Reference(ws, min_col=2, min_row=kpi_rows["Completed Appointments"],
                              max_row=kpi_rows["No Show Appointments"]))
    status.series[0].tx = SeriesLabel(v="Appointments")
    status.series[0].cat = AxDataSource(strLit=StrData(
        ptCount=len(STATUS_LABELS),
        pt=[StrVal(idx=index, v=label) for index, label in enumerate(STATUS_LABELS)]))
    status.dataLabels = value_labels()
    style_chart(status)
    status.width, status.height = STATUS_CHART_WIDTH, CHART_HEIGHT
    ws.add_chart(status, f"A{chart_row + 1 + CHART_ROWS}")

    by_department = BarChart()
    by_department.type = "bar"
    by_department.title = "Appointments by Department (all statuses)"
    total_col = column_index(DEPARTMENT_COLUMNS, "total")
    name_col = column_index(DEPARTMENT_COLUMNS, "department_name")
    by_department.add_data(Reference(department_sheet, min_col=total_col,
                                     min_row=department["header_row"],
                                     max_row=department["last_row"]), titles_from_data=True)
    by_department.set_categories(Reference(department_sheet, min_col=name_col,
                                           min_row=department["first_row"],
                                           max_row=department["last_row"]))
    # First department at the top; the value axis then has to be told to stay at the
    # bottom ("crosses at the maximum category"), away from the title.
    by_department.x_axis.scaling.orientation = "maxMin"
    by_department.y_axis.crosses = "max"
    style_chart(by_department)
    by_department.width, by_department.height = DEPARTMENT_CHART_WIDTH, CHART_HEIGHT
    ws.add_chart(by_department, f"C{chart_row + 1 + CHART_ROWS}")
    return kpi_rows


def value_labels():
    """Data labels that show the value only. Every flag is set explicitly, because
    Excel may treat a missing flag as "show" (that put long names on the old pie)."""
    return DataLabelList(showVal=True, showLegendKey=False, showCatName=False,
                         showSerName=False, showPercent=False, showBubbleSize=False,
                         dLblPos="outEnd")


def column_index(columns, key):
    return next(i for i, c in enumerate(columns, start=1) if c["key"] == key)


def style_chart(chart):
    """One-series bar chart: visible axes, no legend, no axis titles, one blue."""
    chart.x_axis.delete = False   # keep axes visible in current Excel versions
    chart.y_axis.delete = False
    chart.x_axis.title = None
    chart.y_axis.title = None
    chart.y_axis.majorGridlines = None
    chart.legend = None
    for series in chart.series:
        series.graphicalProperties.solidFill = CHART_COLOR
        series.graphicalProperties.line.solidFill = CHART_COLOR


# --- Data Quality sheet -------------------------------------------------------------------

TABLE_DESCRIPTIONS = {
    "departments": "Clinic departments",
    "providers": "Fictional providers",
    "patients": "Fictional patients",
    "provider_schedules": "Provider clinic days (capacity)",
    "appointments": "Booked appointments",
}
VALIDATION_SECTIONS = {
    "duplicate_primary_keys": "Duplicate primary keys",
    "missing_required_values": "Missing required values",
    "foreign_key_violations": "Foreign-key violations",
    "rule_violations": "Rule violations",
}


def write_data_quality(ws, quality, checks):
    write_title(ws, "Data Quality and Reconciliation (Audit)", [
        "Audit sheet: Phase 2 validation results and Phase 3 reconciliation checks. "
        "No other sheet uses values from this sheet.",
    ])
    violations = sum(sum(quality[section].values()) for section in VALIDATION_SECTIONS)
    passed = sum(check["passed"] for check in checks)
    foreign_key = quality["sqlite_foreign_key_check"]

    # Columns C-E hold the count columns (Raw CSV Rows, Processed CSV Rows,
    # Database Rows, Violations, Expected, Actual): wide enough for the header words
    # plus the filter button.
    widths = [34, 80, 22, 22, 22]
    set_widths(ws, widths)
    ws.sheet_view.zoomScale = ANALYSIS_ZOOM

    write_section(ws, 4, "Status")
    status_rows = [
        ["Synthetic data", "Yes: all records are computer-generated"
         if quality["synthetic_data"] else "No"],
        ["Data notice", quality["data_notice"]],
        ["Phase 2 pipeline status", quality["pipeline_status"].capitalize()],
        ["Phase 2 validation violations", violations],
        ["SQLite foreign_key_check",
         f"{'Passed' if foreign_key['passed'] else 'Failed'} "
         f"({foreign_key['violations']} violations)"],
        ["Phase 3 reconciliation checks", f"{passed} of {len(checks)} passed"],
    ]
    row = write_table(ws, 5, ["Item", "Result"], status_rows, [GENERAL, GENERAL],
                      wrap_columns=(2,))
    for status_row in range(6, row + 1):
        ws.cell(row=status_row, column=2).alignment = Alignment(horizontal="left",
                                                                wrap_text=True)
    fit_wrapped_rows(ws, 6, row, 2, widths[1])

    row += 2
    write_section(ws, row, "Row Counts")
    count_rows = [[table, TABLE_DESCRIPTIONS[table], quality["raw_row_counts"][table],
                   quality["processed_csv_row_counts"][table],
                   quality["database_row_counts"][table]] for table in TABLES]
    row = write_table(ws, row + 1, ["Table", "Description", "Raw CSV Rows",
                                    "Processed CSV Rows", "Database Rows"],
                      count_rows, [TEXT, GENERAL, COUNT, COUNT, COUNT], "tblRowCounts")

    row += 2
    write_section(ws, row, "Phase 2 Validation Checks (violations must be 0)")
    check_rows = [[label, name, count]
                  for section, label in VALIDATION_SECTIONS.items()
                  for name, count in quality[section].items()]
    first = row + 2
    row = write_table(ws, row + 1, ["Category", "Check", "Violations"], check_rows,
                      [GENERAL, GENERAL, COUNT], "tblValidationChecks", wrap_columns=(2,))
    fit_wrapped_rows(ws, first, row, 2, widths[1])

    row += 2
    write_section(ws, row, "Phase 3 Reconciliation Checks (breakdowns add up to overall totals)")
    reconciliation_rows = [["PASS" if c["passed"] else "FAIL", c["check"], c["expected"],
                            c["actual"]] for c in checks]
    first = row + 2
    row = write_table(ws, row + 1, ["Result", "Check", "Expected", "Actual"],
                      reconciliation_rows, [GENERAL, GENERAL, GENERAL, GENERAL],
                      "tblReconciliation", wrap_columns=(2,))
    fit_wrapped_rows(ws, first, row, 2, widths[1])


# --- Data sheets ----------------------------------------------------------------------------

APPOINTMENT_FORMATS = {
    "appointment_id": TEXT, "patient_id": TEXT, "provider_id": TEXT, "department_id": TEXT,
    "appointment_date": DATE, "scheduled_datetime": DATETIME, "check_in_datetime": DATETIME,
    "seen_datetime": DATETIME, "checkout_datetime": DATETIME, "wait_minutes": MINUTES,
    "visit_minutes": MINUTES, "appointment_type": TEXT, "appointment_status": TEXT,
}
CAPACITY_FORMATS = {
    "schedule_date": DATE, "provider_id": TEXT, "provider_name": TEXT, "department_id": TEXT,
    "department_name": TEXT, "daily_slot_capacity": COUNT, "available_slots": COUNT,
}
CAPACITY_QUERY = """
    SELECT s.schedule_date, s.provider_id, p.provider_name, s.department_id,
           d.department_name, p.daily_slot_capacity, s.available_slots
    FROM provider_schedules AS s
    JOIN providers AS p ON p.provider_id = s.provider_id
    JOIN departments AS d ON d.department_id = s.department_id
    ORDER BY s.schedule_date, s.provider_id
"""


def to_excel_values(df, formats):
    """ISO date/datetime text -> Python date/datetime, so Excel stores real dates."""
    converted = df.astype(object).where(df.notna(), None)
    for column, number_format in formats.items():
        if number_format == DATE:
            converted[column] = [date.fromisoformat(v) if v else None for v in df[column]]
        elif number_format == DATETIME:
            converted[column] = [datetime.fromisoformat(v) if isinstance(v, str) else None
                                 for v in df[column]]
    return converted.values.tolist()


def data_sheet_width(df, column, number_format):
    """Wide enough for the header plus its filter button, and for the values."""
    minimum = {DATETIME: 18, DATE: 12}.get(number_format, 10)
    longest_text = df[column].astype(str).str.len().max() if number_format == TEXT else 0
    return max(minimum, len(column) + 5, longest_text + 3)


def write_data_sheet(ws, df, formats, table_name):
    columns = list(df.columns)
    write_table(ws, 1, columns, to_excel_values(df, formats),
                [formats[c] for c in columns], table_name)
    ws.freeze_panes = "A2"
    set_widths(ws, [data_sheet_width(df, c, formats[c]) for c in columns])


# --- Putting it together ----------------------------------------------------------------------

def read_database(database_path):
    with closing(sqlite3.connect(database_path)) as connection:
        appointments = pd.read_sql_query(
            "SELECT * FROM appointments ORDER BY appointment_id", connection)
        capacity = pd.read_sql_query(CAPACITY_QUERY, connection)
        row_counts = {table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                      for table in TABLES}
    return appointments, capacity, row_counts


def read_quality_report(quality_report_path, database_row_counts):
    path = Path(quality_report_path)
    if not path.exists():
        raise FileNotFoundError(f"Data-quality report not found: {path}\n"
                                "Create it by running:\n    python prepare_data.py")
    quality = json.loads(path.read_text(encoding="utf-8"))
    if quality.get("pipeline_status") != "passed":
        raise ReportError(f"{path} does not show a passed pipeline. Fix the data and run "
                          "python prepare_data.py again.")
    if quality.get("database_row_counts") != database_row_counts:
        raise ReportError(f"{path} does not match the database row counts. "
                          "Run python prepare_data.py again.")
    return quality


def build_report(database_path=config.DATABASE_FILE,
                 quality_report_path=config.QUALITY_REPORT_FILE,
                 output_path=config.EXCEL_REPORT_FILE):
    results, checks = analyze_data.load_analytics(database_path)
    appointments, capacity, row_counts = read_database(database_path)
    quality = read_quality_report(quality_report_path, row_counts)

    workbook = Workbook()
    sheets = {name: workbook.create_sheet(name) for name in SHEET_NAMES}
    workbook.remove(workbook["Sheet"])

    monthly_df = results["monthly_trends"].copy()
    monthly_df["month_start"] = [date(int(m[:4]), int(m[5:7]), 1) for m in monthly_df["month"]]
    monthly = write_analysis_sheet(
        sheets["Monthly Trends"], "Monthly Trends",
        ["One row per month, in date order. Rates are formulas (numerator / denominator, "
         "blank when the denominator is 0)."],
        monthly_df, MONTHLY_COLUMNS, "tblMonthlyTrends", "B5")
    department = write_analysis_sheet(
        sheets["Department Analysis"], "Department Analysis",
        ["One row per department, sorted by department ID (the order the departments "
         "are defined in config.py)."],
        results["department_performance"], DEPARTMENT_COLUMNS, "tblDepartments", "C5")
    write_analysis_sheet(
        sheets["Provider Analysis"], "Provider Analysis",
        ["One row per provider, sorted by department ID and then provider ID. Available "
         "Slots is the sum of the provider's scheduled clinic-day slots."],
        results["provider_performance"], PROVIDER_COLUMNS, "tblProviders", "C5")
    write_analysis_sheet(
        sheets["Appointment Types"], "Appointment Types",
        ["Sorted by total appointments (highest first). No capacity metrics: slots are not "
         "reserved per appointment type, so there is no valid capacity denominator."],
        results["appointment_type_performance"], TYPE_COLUMNS, "tblAppointmentTypes", "B5")
    write_summary(sheets["Summary"], results["overall_kpis"].iloc[0], monthly,
                  sheets["Monthly Trends"], department, sheets["Department Analysis"])
    write_data_quality(sheets["Data Quality"], quality, checks)
    write_data_sheet(sheets["Appointment Data"], appointments, APPOINTMENT_FORMATS,
                     "tblAppointmentData")
    write_data_sheet(sheets["Capacity Data"], capacity, CAPACITY_FORMATS, "tblCapacityData")

    workbook.properties.title = "Healthcare Patient Flow & Operational Analytics Report"
    workbook.properties.subject = "Synthetic data only"
    # Neutral author fields: no personal names, email addresses or machine details.
    workbook.properties.creator = WORKBOOK_AUTHOR
    workbook.properties.lastModifiedBy = WORKBOOK_AUTHOR
    # Ask Excel to recalculate every formula when the workbook is opened.
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    workbook.active = 0

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main(argv=None):
    parser = argparse.ArgumentParser(description="Create the Excel report (synthetic data).")
    parser.add_argument("--database", default=str(config.DATABASE_FILE),
                        help="SQLite database (default data/processed/healthcare_analytics.db)")
    parser.add_argument("--quality-report", default=str(config.QUALITY_REPORT_FILE),
                        help="Phase 2 report (default data/output/phase2_data_quality.json)")
    parser.add_argument("--output", default=str(config.EXCEL_REPORT_FILE),
                        help="workbook path (default reports/healthcare_patient_flow_report.xlsx)")
    args = parser.parse_args(argv)

    try:
        output_path = build_report(args.database, args.quality_report, args.output)
    except PermissionError:
        print(f"ERROR: cannot write {args.output}. Close the workbook in Excel and try again.",
              file=sys.stderr)
        return 1
    except (FileNotFoundError, ReportError, analyze_data.AnalyticsError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print("Excel report created (SYNTHETIC data).")
    print(f"  {output_path.resolve()}  ({output_path.stat().st_size / 1024:,.0f} KB)")
    print(f"  Sheets: {', '.join(SHEET_NAMES)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
